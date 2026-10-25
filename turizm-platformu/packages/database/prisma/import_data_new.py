#!/usr/bin/env python3
"""
Yeni Prisma Şeması İçin Veri İçe Aktarma Scripti
ödev/ klasöründeki tüm destinasyon verilerini yeni normalize edilmiş şemaya aktarır.
"""

import os
import re
import sys
import shutil
import uuid
from pathlib import Path
from decimal import Decimal
from typing import Optional, List, Tuple, Dict, Any
from datetime import datetime

import pg8000.dbapi as pgdb
from docx import Document
import openpyxl
import pandas as pd

DB_CONFIG = {
    "host": "localhost", "port": 5432, "database": "trabzon_turizm",
    "user": "postgres", "password": "1234",
}

ILLER_DATA = [
    {"ad": "Trabzon", "slug": "trabzon", "code": "61", "lat": 41.0015, "lon": 39.7178, "zoom": 9},
    {"ad": "Giresun", "slug": "giresun", "code": "28", "lat": 40.9128, "lon": 38.3895, "zoom": 9},
    {"ad": "Ordu", "slug": "ordu", "code": "52", "lat": 40.9839, "lon": 37.8764, "zoom": 9},
    {"ad": "Gümüşhane", "slug": "gumushane", "code": "29", "lat": 40.4386, "lon": 39.4803, "zoom": 9},
    {"ad": "Rize", "slug": "rize", "code": "53", "lat": 41.0201, "lon": 40.5216, "zoom": 9},
    {"ad": "Artvin", "slug": "artvin", "code": "08", "lat": 41.1828, "lon": 41.8183, "zoom": 9},
]

KATEGORILER = [
    {"ad": "Cami", "slug": "cami", "icon": "mosque", "color": "#1B5E20", "sortOrder": 1},
    {"ad": "Türbe", "slug": "turbe", "icon": "tomb", "color": "#4E342E", "sortOrder": 2},
    {"ad": "Kilise", "slug": "kilise", "icon": "church", "color": "#C62828", "sortOrder": 3},
    {"ad": "Şapel", "slug": "sapel", "icon": "chapel", "color": "#AD1457", "sortOrder": 4},
    {"ad": "Kale", "slug": "kale", "icon": "castle", "color": "#BF360C", "sortOrder": 5},
    {"ad": "Tarihi Bina", "slug": "tarihi-bina", "icon": "building", "color": "#5D4037", "sortOrder": 6},
    {"ad": "Han", "slug": "han", "icon": "warehouse", "color": "#795548", "sortOrder": 7},
    {"ad": "Hamam", "slug": "hamam", "icon": "spa", "color": "#00695C", "sortOrder": 8},
    {"ad": "Şelale", "slug": "selale", "icon": "waterfall", "color": "#0277BD", "sortOrder": 9},
    {"ad": "Köprü", "slug": "kopru", "icon": "bridge", "color": "#546E7A", "sortOrder": 10},
    {"ad": "Sivil Mimari", "slug": "sivil-mimari", "icon": "architecture", "color": "#37474F", "sortOrder": 11},
    {"ad": "Mağara", "slug": "magara", "icon": "cave", "color": "#424242", "sortOrder": 12},
    {"ad": "Milli Park", "slug": "milli-park", "icon": "park", "color": "#2E7D32", "sortOrder": 13},
    {"ad": "Müze", "slug": "muze", "icon": "museum", "color": "#6A1B9A", "sortOrder": 14},
    {"ad": "Şenlikler", "slug": "senlikler", "icon": "festival", "color": "#E65100", "sortOrder": 15},
    {"ad": "Göl", "slug": "gol", "icon": "lake", "color": "#0288D1", "sortOrder": 16},
    {"ad": "Kanyon", "slug": "kanyon", "icon": "terrain", "color": "#3E2723", "sortOrder": 17},
    {"ad": "Yerleşim", "slug": "yerlesim", "icon": "location_city", "sortOrder": 18},
    {"ad": "Çeşme", "slug": "cesme", "icon": "fountain", "sortOrder": 19},
    {"ad": "Değirmen", "slug": "debirmen", "icon": "grain", "sortOrder": 20},
    {"ad": "Köy", "slug": "koy", "icon": "home", "sortOrder": 21},
    {"ad": "Antik", "slug": "antik", "icon": "history_edu", "sortOrder": 22},
    {"ad": "Anıt", "slug": "anit", "icon": "monument", "sortOrder": 23},
    {"ad": "Kültür Evi", "slug": "kultur-evi", "icon": "account_balance", "sortOrder": 24},
    {"ad": "Medrese", "slug": "medrese", "icon": "school", "sortOrder": 25},
    {"ad": "Kaymakamlık", "slug": "kaymakamlik", "icon": "account_balance", "sortOrder": 26},
    {"ad": "Okul", "slug": "okul", "icon": "school", "sortOrder": 27},
    {"ad": "Bahçe Kapısı", "slug": "bahce-kapisi", "icon": "yard", "sortOrder": 28},
    {"ad": "Kaya Mezarı", "slug": "kaya-mezari", "icon": "bedtime", "sortOrder": 29},
    {"ad": "Kayaüstü", "slug": "kayaustu", "icon": "landscape", "sortOrder": 30},
    {"ad": "Kayalık", "slug": "kayalik", "icon": "terrain", "sortOrder": 32},
    {"ad": "Toprak Altı", "slug": "toprak-ali", "icon": "layers", "sortOrder": 33},
    {"ad": "Yamaç", "slug": "yamac", "icon": "terrain", "sortOrder": 33},
    {"ad": "Arkeolojik", "slug": "arkeolojik", "icon": "history_edu", "sortOrder": 34},
]

KLASOR_TUR_ESLESTIRME = {
    "camiler-türbeler": "Cami", "türbeler": "Türbe", "kilise-manastır-şapel": "Kilise",
    "kale": "Kale", "göl": "Göl", "şelale": "Şelale", "müzeler": "Müze",
    "milli park": "Milli Park", "kanyon": "Kanyon", "mağara": "Mağara",
    "köprüler": "Köprü", "han-hamam": "Hamam", "sivil mimari": "Sivil Mimari",
    "yerleşim": "Yerleşim", "şenlikler": "Şenlikler", "yemek listesi": None,
    "camii": "Cami", "değirmen": "Değirmen", "göller": "Göl", "hamam": "Hamam",
    "kale": "Kale", "kanyonlar": "Kanyon", "kilise-manastır-şapel": "Kilise",
    "köprü": "Köprü", "mağara": "Mağara", "milli park": "Milli Park",
    "sivil mimari": "Sivil Mimari", "türbe": "Türbe", "yerleşim": "Yerleşim",
    "şelale": "Şelale", "şenlikler": "Şenlikler",
    "camii": "Cami", "göl": "Göl", "kale": "Kale", "köprü": "Köprü",
    "mağara": "Mağara", "milli park": "Milli Park", "sivil mimari": "Sivil Mimari",
    "türbe - mezarlık": "Türbe", "yerleşim": "Yerleşim", "şelale": "Şelale",
    "cami - türbe": "Cami", "göl": "Göl", "hamam": "Hamam", "kale": "Kale",
    "kanyon": "Kanyon", "kilise": "Kilise", "köprü": "Köprü",
    "sivil mimari": "Sivil Mimari", "tarih": "Tarihi Bina", "yerleşim": "Yerleşim",
    "çeşme - su kuyusu": "Çeşme",
    "göl": "Göl", "şelale": "Şelale", "milli park": "Milli Park",
}

ARTVIN_TUR_IPUCLARI = {
    "cami": "Cami", "türbe": "Türbe", "kalesi": "Kale", "kilise": "Kilise",
    "manastır": "Kilise", "şelale": "Şelale", "mezarlık": "Kaya Mezarı",
    "hamam": "Hamam", "köprü": "Köprü", "konut": "Sivil Mimari",
    "müze": "Müze", "kültür evi": "Kültür Evi", "medrese": "Medrese",
    "kaymakamlık": "Kaymakamlık", "lohman": "Tarihi Bina", "okul": "Okul",
    "anıtı": "Anıt", "bahçe kapısı": "Bahçe Kapısı", "kaya mezarı": "Kaya Mezarı",
    "kayaüstü": "Kayaüstü", "kayalık": "Kayalık", "toprak altı": "Toprak Altı",
    "antik": "Antik", "yerleşim": "Yerleşim", "arkeolojik": "Arkeolojik", "yamaç": "Yamaç",
}


def get_db_connection():
    return pgdb.connect(
        user=DB_CONFIG["user"], password=DB_CONFIG["password"] or None,
        host=DB_CONFIG["host"], port=int(DB_CONFIG["port"]),
        database=DB_CONFIG["database"], timeout=10,
        application_name="TurizmVeriImportNew",
    )


def slugify(text: str) -> str:
    text = text.lower().strip()
    tr_map = str.maketrans({'ç':'c','ğ':'g','ı':'i','ö':'o','ş':'s','ü':'u','Ç':'c','Ğ':'g','İ':'i','Ö':'o','Ş':'s','Ü':'u'})
    text = text.translate(tr_map)
    text = re.sub(r'[^a-z0-9\s-]', '', text)
    text = re.sub(r'\s+', '-', text)
    text = re.sub(r'-+', '-', text)
    return text.strip('-')


def ensure_unique_slug(conn, table: str, base_slug: str, exclude_id: str = None) -> str:
    slug = base_slug
    counter = 1
    while True:
        cur = conn.cursor()
        if exclude_id:
            cur.execute(f'SELECT 1 FROM {table} WHERE slug = %s AND id != %s;', (slug, exclude_id))
        else:
            cur.execute(f'SELECT 1 FROM {table} WHERE slug = %s;', (slug,))
        exists = cur.fetchone()
        cur.close()
        if not exists:
            return slug
        counter += 1
        slug = f"{base_slug}-{counter}"


def read_docx_text(filepath: Path) -> str:
    try:
        doc = Document(filepath)
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        seen = set(parts)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text and text not in seen:
                        parts.append(text)
                        seen.add(text)
        return "\n".join(parts).replace("\x00", "")
    except Exception as e:
        print(f"  [UYARI] DOCX okuma hatasi ({filepath.name}): {e}")
        return ""


def parse_coordinates(text: str) -> Tuple[Optional[Decimal], Optional[Decimal]]:
    text = text.strip().replace(",", ".").replace(";", ",")
    numbers = re.findall(r'-?\d+\.?\d*', text)
    if len(numbers) >= 2:
        lat = Decimal(numbers[0]); lon = Decimal(numbers[1])
        if -90 <= lat <= 90 and -180 <= lon <= 180:
            return lat.quantize(Decimal("0.000001")), lon.quantize(Decimal("0.000001"))
    return None, None


def find_image_file(dest_folder: Path) -> Optional[str]:
    for ext in ['.jpg', '.jpeg', '.png', '.jfif']:
        images = list(dest_folder.glob(f'*{ext}'))
        if images: return images[0].name
    return None


def normalize_name(name: str) -> str:
    name = re.sub(r'\s*\([^)]+\)\s*$', '', name).strip()
    return re.sub(r'\s+', ' ', name).strip()


def determine_type_from_folder(folder_name: str, province: str) -> Optional[str]:
    folder_lower = folder_name.lower().strip()
    if folder_lower in KLASOR_TUR_ESLESTIRME:
        return KLASOR_TUR_ESLESTIRME[folder_lower]
    if province == "Artvin":
        for keyword, tur in ARTVIN_TUR_IPUCLARI.items():
            if keyword in folder_lower:
                return tur
    for keyword, tur in KLASOR_TUR_ESLESTIRME.items():
        if keyword and keyword in folder_lower:
            return tur
    return None


def get_db_connection():
    return pgdb.connect(
        user=DB_CONFIG["user"], password=DB_CONFIG["password"] or None,
        host=DB_CONFIG["host"], port=int(DB_CONFIG["port"]),
        database=DB_CONFIG["database"], timeout=10,
        application_name="TurizmVeriImportNew",
    )


def setup_provinces_and_categories(conn):
    cur = conn.cursor()
    try:
        for il in ILLER_DATA:
            pid = str(uuid.uuid4())
            cur.execute("""
                INSERT INTO provinces (id, name, slug, code, latitude, longitude, zoom, "createdAt", "updatedAt")
                VALUES (%s, %s, %s, %s, %s, %s, %s, NOW(), NOW())
                ON CONFLICT (slug) DO UPDATE SET
                    name = EXCLUDED.name, code = EXCLUDED.code,
                    latitude = EXCLUDED.latitude, longitude = EXCLUDED.longitude,
                    zoom = EXCLUDED.zoom, "updatedAt" = NOW()
                RETURNING id;
            """, (pid, il["ad"], il["slug"], il["code"], il["lat"], il["lon"], il["zoom"]))
            row = cur.fetchone()
            if row: print(f"  [Province] {il['ad']} (ID: {row[0]})")

        for kat in KATEGORILER:
            cid = str(uuid.uuid4())
            color = kat.get("color", "#37474F")
            cur.execute("""
                INSERT INTO destination_types (id, name, slug, icon, color, description, "sortOrder", "isActive", "createdAt", "updatedAt")
                VALUES (%s, %s, %s, %s, %s, %s, %s, true, NOW(), NOW())
                ON CONFLICT (slug) DO UPDATE SET
                    name = EXCLUDED.name, icon = EXCLUDED.icon, color = EXCLUDED.color,
                    description = EXCLUDED.description, "sortOrder" = EXCLUDED."sortOrder", "updatedAt" = NOW()
                RETURNING id;
            """, (cid, kat["ad"], kat["slug"], kat["icon"], color, f"{kat['ad']} türü destinasyonları", kat["sortOrder"]))
            row = cur.fetchone()
            if row: print(f"  [Category] {kat['ad']} (ID: {row[0]})")

        conn.commit()
        print("[OK] Provinces ve Categories hazır.")
    finally:
        cur.close()


def get_province_id(conn, province_name: str) -> str:
    cur = conn.cursor()
    try:
        cur.execute('SELECT id FROM provinces WHERE lower(name) = lower(%s);', (province_name,))
        row = cur.fetchone()
        if row: return row[0]
        raise ValueError(f"Province bulunamadı: {province_name}")
    finally: cur.close()


def get_category_id(conn, category_name: str) -> str:
    cur = conn.cursor()
    try:
        cur.execute('SELECT id FROM destination_types WHERE lower(name) = lower(%s);', (category_name,))
        row = cur.fetchone()
        if row: return row[0]
        cur.execute('SELECT id FROM destination_types WHERE lower(name) = lower(%s);', ("Tarihi Bina",))
        row = cur.fetchone()
        if row: return row[0]
        raise ValueError(f"Category bulunamadı: {category_name}")
    finally: cur.close()


def get_district_id(conn, province_id: str, district_name: str) -> Optional[str]:
    cur = conn.cursor()
    try:
        cur.execute('SELECT id FROM districts WHERE "provinceId" = %s AND lower(name) = lower(%s);', (province_id, district_name))
        row = cur.fetchone()
        return row[0] if row else None
    finally: cur.close()


IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".jfif"}

# id dahil 25 kolon (createdAt/updatedAt dahil)
DEST_COLS = "".join((
    '"id", "name", "slug", "description", "shortDescription", "typeId", "provinceId", "districtId",'
    '"latitude", "longitude", "address", "phone", "website", "openingHours", "images", "thumbnailUrl",'
    '"metadata", "isActive", "isFeatured", "viewCount", "clickCount", "averageRating", "reviewCount",'
    '"createdAt", "updatedAt"'
))

DEST_INSERT = f"""
    INSERT INTO destinations ({DEST_COLS})
    VALUES ({', '.join(['%s'] * 25)})
    ON CONFLICT ("slug") DO UPDATE SET
        "name" = EXCLUDED."name",
        "description" = COALESCE(EXCLUDED."description", destinations."description"),
        "shortDescription" = COALESCE(EXCLUDED."shortDescription", destinations."shortDescription"),
        "images" = CASE
            WHEN cardinality(EXCLUDED."images") > 0 THEN EXCLUDED."images"
            ELSE destinations."images"
        END,
        "thumbnailUrl" = COALESCE(EXCLUDED."thumbnailUrl", destinations."thumbnailUrl"),
        "updatedAt" = NOW()
    RETURNING id;
"""

DEST_TRANSL_INSERT = """
    INSERT INTO destination_translations ("id", "destinationId", "langCode", "name", "description", "shortDescription")
    VALUES (gen_random_uuid(), %s, %s, %s, %s)
    ON CONFLICT ("destinationId", "langCode") DO UPDATE SET
        "name" = EXCLUDED."name", "description" = EXCLUDED."description",
        "shortDescription" = EXCLUDED."shortDescription";
"""


def clean_text(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    text = value.replace("\x00", "").strip()
    return text or None


def list_image_files(dest_folder: Path) -> List[Path]:
    files: List[Path] = []
    try:
        entries = list(dest_folder.iterdir())
    except OSError:
        return files
    for entry in entries:
        if not entry.is_file() or entry.name.startswith("~$"):
            continue
        if entry.suffix.lower() in IMAGE_SUFFIXES:
            files.append(entry)
    return sorted(files, key=lambda item: item.name.lower())


def is_coordinate_name(name: str) -> bool:
    lowered = name.lower()
    return lowered == "k.docx" or (lowered.startswith("koord") and lowered.endswith((".docx", ".txt")))


def is_description_name(name: str) -> bool:
    lowered = name.lower()
    return lowered in {"m.docx", "metin.txt"} or (lowered.startswith("metin") and lowered.endswith(".docx"))


def is_destination_folder(folder: Path) -> bool:
    """Koordinat dosyasi varsa bu klasor bir destinasyondur, tur klasoru degil."""
    try:
        names = {item.name.lower() for item in folder.iterdir() if item.is_file() and not item.name.startswith("~$")}
    except OSError:
        return False
    return any(is_coordinate_name(name) for name in names)


def read_source_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return read_docx_text(path)
    if suffix == ".txt":
        for encoding in ("utf-8", "windows-1254", "cp1254"):
            try:
                return path.read_text(encoding=encoding).replace("\x00", "").strip()
            except UnicodeDecodeError:
                continue
    return ""


def first_matching_file(folder: Path, predicate) -> Optional[Path]:
    preferred_names = {"metin.docx", "m.docx", "koordinat.docx", "k.docx"}
    try:
        files = [item for item in folder.iterdir() if item.is_file() and not item.name.startswith("~$")]
    except OSError:
        return None
    matches = [item for item in files if predicate(item.name.lower())]
    if not matches:
        return None
    preferred = [item for item in matches if item.name.lower() in preferred_names]
    chosen = preferred or matches
    return sorted(chosen, key=lambda item: item.name.lower())[0]


def register_images(dest_folder: Path, images_dest: Path) -> List[str]:
    """Klasordeki fotograflari images/ altina kopyalar, goreli yollari dondurur."""
    images_dest.mkdir(parents=True, exist_ok=True)
    relative_paths: List[str] = []
    for src in list_image_files(dest_folder):
        dst = images_dest / src.name
        if not dst.exists():
            try:
                shutil.copy2(src, dst)
            except OSError as exc:
                print(f"  [UYARI] Gorsel kopyalanamadi ({src.name}): {exc}")
                continue
        relative_paths.append(f"images/{src.name}")
    return relative_paths


def find_destination_ids_at(conn, province_id: str, enlem, boylam) -> List[str]:
    cur = conn.cursor()
    try:
        cur.execute(
            """
            SELECT id
              FROM destinations
             WHERE "provinceId" = %s
               AND latitude = %s
               AND longitude = %s
             ORDER BY "createdAt" ASC, id ASC;
            """,
            (province_id, enlem, boylam),
        )
        return [row[0] for row in cur.fetchall()]
    finally:
        cur.close()


def find_existing_destination(conn, province_id: str, name: str, enlem, boylam) -> Optional[str]:
    """Ayni il, ad ve koordinattaki kaydi bulur. Numarasiz slug tercih edilir."""
    cur = conn.cursor()
    try:
        cur.execute(
            """
            SELECT id
              FROM destinations
             WHERE "provinceId" = %s
               AND lower(btrim(name)) = lower(btrim(%s))
               AND latitude = %s
               AND longitude = %s
             ORDER BY CASE WHEN slug ~ '-[0-9]+$' THEN 1 ELSE 0 END,
                      "createdAt" ASC,
                      id ASC
             LIMIT 1;
            """,
            (province_id, name, enlem, boylam),
        )
        row = cur.fetchone()
        return row[0] if row else None
    finally:
        cur.close()


def coordinates_taken(conn, province_id: str, enlem, boylam) -> bool:
    cur = conn.cursor()
    try:
        cur.execute(
            """
            SELECT 1
              FROM destinations
             WHERE "provinceId" = %s
               AND latitude = %s
               AND longitude = %s
             LIMIT 1;
            """,
            (province_id, enlem, boylam),
        )
        return cur.fetchone() is not None
    finally:
        cur.close()


def upsert_translation(cur, dest_id: str, ad: str, aciklama: Optional[str]) -> None:
    text = clean_text(aciklama)
    short = text[:500] if text else None
    cur.execute(
        """
        INSERT INTO destination_translations
            ("id", "destinationId", "langCode", "name", "description", "shortDescription")
        VALUES (gen_random_uuid(), %s, 'tr', %s, %s, %s)
        ON CONFLICT ("destinationId", "langCode") DO UPDATE SET
            "name" = EXCLUDED."name",
            "description" = COALESCE(EXCLUDED."description", destination_translations."description"),
            "shortDescription" = COALESCE(EXCLUDED."shortDescription", destination_translations."shortDescription");
        """,
        (dest_id, ad, text, short),
    )


def make_dest_values(ad, slug, aciklama, tur_id, province_id, enlem, boylam, image_paths):
    """25 parametreli values tuple'i olusturur (id + createdAt/updatedAt dahil)."""
    if isinstance(image_paths, str):
        paths = [image_paths] if image_paths else []
    else:
        paths = list(image_paths or [])
    text = clean_text(aciklama)
    short = text[:500] if text else None
    thumb = paths[0] if paths else None
    return (
        str(uuid.uuid4()),
        ad, slug, text, short,
        tur_id, province_id, None,
        enlem, boylam, None, None, None,
        "{}",
        paths, thumb, "{}",
        True, False, 0, 0, Decimal("0.00"), 0,
        datetime.now(), datetime.now(),
    )


def save_destination(conn, ad, aciklama, tur_id, province_id, enlem, boylam, image_paths, update_content: bool) -> Optional[str]:
    """Ayni il + ad + koordinat varsa gunceller, yoksa tek satir ekler."""
    text = clean_text(aciklama)
    short = text[:500] if text else None
    paths = list(image_paths or [])
    thumb = paths[0] if paths else None
    existing_id = find_existing_destination(conn, province_id, ad, enlem, boylam)
    if not existing_id:
        at_coords = find_destination_ids_at(conn, province_id, enlem, boylam)
        if len(at_coords) == 1:
            existing_id = at_coords[0]
        elif len(at_coords) > 1:
            print(f"  [UYARI] Ayni koordinatta birden fazla kayit var, atlandi: {ad}")
            return None

    cur = conn.cursor()
    try:
        if existing_id:
            if update_content:
                cur.execute(
                    """
                    UPDATE destinations
                       SET description = COALESCE(%s, description),
                           "shortDescription" = COALESCE(%s, "shortDescription"),
                           images = CASE
                               WHEN cardinality(%s::text[]) > 0 THEN %s::text[]
                               ELSE images
                           END,
                           "thumbnailUrl" = COALESCE(%s, "thumbnailUrl"),
                           "updatedAt" = NOW()
                     WHERE id = %s;
                    """,
                    (text, short, paths, paths, thumb, existing_id),
                )
                upsert_translation(cur, existing_id, ad, text)
                conn.commit()
            return existing_id

        slug = ensure_unique_slug(conn, "destinations", slugify(ad) or "destinasyon")
        vals = make_dest_values(ad, slug, text, tur_id, province_id, enlem, boylam, paths)
        cur.execute(DEST_INSERT, vals)
        row = cur.fetchone()
        if not row:
            conn.rollback()
            return None
        dest_id = row[0]
        upsert_translation(cur, dest_id, ad, text)
        conn.commit()
        return dest_id
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()


def process_destination_folder_new(conn, dest_folder: Path, province_name: str, province_id: str, folder_type: Optional[str], images_dest: Path) -> int:
    ad = normalize_name(dest_folder.name)
    metin_file = first_matching_file(dest_folder, is_description_name)
    koordinat_file = first_matching_file(dest_folder, is_coordinate_name)
    aciklama = read_source_text(metin_file) if metin_file else ""
    enlem, boylam = parse_coordinates(read_source_text(koordinat_file)) if koordinat_file else (None, None)

    if enlem is None or boylam is None:
        print(f"  [UYARI] Koordinat bulunamadi: {ad}")
        return 0

    tur_adi = (
        folder_type
        or determine_type_from_folder(dest_folder.name, province_name)
        or determine_type_from_folder(dest_folder.parent.name, province_name)
        or "Tarihi Bina"
    )
    image_paths = register_images(dest_folder, images_dest)
    existed = find_existing_destination(conn, province_id, ad, enlem, boylam) is not None
    if not existed:
        existed = len(find_destination_ids_at(conn, province_id, enlem, boylam)) >= 1
    try:
        tur_id = get_category_id(conn, tur_adi)
        saved = save_destination(
            conn, ad, aciklama, tur_id, province_id, enlem, boylam, image_paths, update_content=True,
        )
    except Exception as exc:
        print(f"  [HATA] Kayit hatasi ({ad}): {exc}")
        return 0
    if not saved:
        return 0
    action = "GUNCELLENDI" if existed else "EKLENDI"
    print(f"  [{action}] {ad} - [{float(enlem):.6f}, {float(boylam):.6f}] ({len(image_paths)} resim)")
    return 1


def process_csv_file_new(conn, csv_path: Path, province_name: str, province_id: str, default_type: str, images_dest: Path) -> int:
    count = 0
    for enc in ['utf-8', 'windows-1254', 'iso-8859-9', 'cp1254']:
        try:
            df = pd.read_csv(csv_path, sep=';', encoding=enc)
            break
        except UnicodeDecodeError: continue
        except Exception as e:
            if 'codec' not in str(e).lower(): raise
    else:
        print(f"  [UYARI] CSV okunamadi: {csv_path.name}")
        return 0

    try:
        df.columns = [c.strip().lower() for c in df.columns]
        tur_id = get_category_id(conn, default_type)
        for _, row in df.iterrows():
            ad = str(row.get('adı', '') or row.get('ad', '') or '').strip()
            if not ad or ad.lower() == 'nan': continue
            enlem = row.get('enlem', row.get('lat', None))
            boylam = row.get('boylam', row.get('lon', row.get('boylam', None)))
            if pd.isna(enlem) or pd.isna(boylam): continue
            enlem = Decimal(str(enlem)).quantize(Decimal("0.000001"))
            boylam = Decimal(str(boylam)).quantize(Decimal("0.000001"))
            if find_existing_destination(conn, province_id, ad, enlem, boylam) or coordinates_taken(conn, province_id, enlem, boylam):
                continue
            try:
                saved = save_destination(
                    conn, ad, None, tur_id, province_id, enlem, boylam, [], update_content=False,
                )
                if saved:
                    count += 1
            except Exception as e:
                print(f"  [UYARI] CSV kayit hatasi ({ad}): {e}")
    except Exception as e:
        print(f"  [UYARI] CSV isleme hatasi ({csv_path.name}): {e}")
    return count


def process_excel_file_new(conn, excel_path: Path, province_name: str, province_id: str, default_type: str, images_dest: Path) -> int:
    count = 0
    try:
        wb = openpyxl.load_workbook(excel_path)
        tur_id = get_category_id(conn, default_type)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows: continue
            for row_data in rows[1:]:
                if not any(row_data): continue
                ad = str(row_data[0] or '').strip() if len(row_data) > 0 else ''
                if not ad or ad.lower() == 'none': continue
                enlem = row_data[1] if len(row_data) > 1 else None
                boylam = row_data[2] if len(row_data) > 2 else None
                if enlem is None or boylam is None: continue
                try:
                    enlem = Decimal(str(enlem)).quantize(Decimal("0.000001"))
                    boylam = Decimal(str(boylam)).quantize(Decimal("0.000001"))
                except Exception:
                    continue
                if find_existing_destination(conn, province_id, ad, enlem, boylam) or coordinates_taken(conn, province_id, enlem, boylam):
                    continue
                try:
                    saved = save_destination(
                        conn, ad, None, tur_id, province_id, enlem, boylam, [], update_content=False,
                    )
                    if saved:
                        count += 1
                except Exception as e:
                    print(f"  [UYARI] Excel kayit hatasi ({ad}): {e}")
    except Exception as e:
        print(f"  [UYARI] Excel isleme hatasi ({excel_path.name}): {e}")
    return count


def process_tree(conn, folder: Path, province_name: str, province_id: str, inherited_type: Optional[str], images_dest: Path) -> int:
    """Tur klasorlerini dolasir. koordinat dosyasi olan klasor bir kez yazilir."""
    total = 0
    here_type = determine_type_from_folder(folder.name, province_name) or inherited_type
    default_type = here_type or "Tarihi Bina"

    for csv_file in sorted(folder.glob("*.csv")):
        if csv_file.name.startswith("~$"):
            continue
        print(f"  CSV: {csv_file.name} -> {default_type}")
        total += process_csv_file_new(conn, csv_file, province_name, province_id, default_type, images_dest)
    for excel_file in sorted(folder.glob("*.xlsx")):
        if excel_file.name.startswith("~$"):
            continue
        print(f"  Excel: {excel_file.name} -> {default_type}")
        total += process_excel_file_new(conn, excel_file, province_name, province_id, default_type, images_dest)

    try:
        children = [item for item in folder.iterdir() if item.is_dir() and not item.name.startswith("~$")]
    except OSError as exc:
        print(f"  [UYARI] Klasor okunamadi ({folder.name}): {exc}")
        return total

    for child in sorted(children, key=lambda item: item.name.lower()):
        if is_destination_folder(child):
            total += process_destination_folder_new(
                conn, child, province_name, province_id, here_type, images_dest,
            )
        else:
            total += process_tree(conn, child, province_name, province_id, here_type, images_dest)
    return total


def process_province_new(conn, province_path: Path, province_name: str, province_id: str, images_dest: Path) -> int:
    print(f"\n{'='*60}\nIsleniyor: {province_name}\n{'='*60}")
    total = process_tree(conn, province_path, province_name, province_id, None, images_dest)
    print(f"\n{province_name} islenen: {total}")
    return total


def dedupe_destinations(conn) -> int:
    """Ayni il, ad ve koordinattaki kopyalari siler. Numarasiz slug kalir."""
    cur = conn.cursor()
    try:
        cur.execute("SELECT COUNT(*) FROM destinations;")
        before = int(cur.fetchone()[0])
        cur.execute(
            """
            WITH ranked AS (
                SELECT id,
                       ROW_NUMBER() OVER (
                           PARTITION BY "provinceId", lower(btrim(name)), latitude, longitude
                           ORDER BY
                               CASE WHEN slug ~ '-[0-9]+$' THEN 1 ELSE 0 END,
                               "createdAt" ASC,
                               id ASC
                       ) AS rn
                FROM destinations
            ),
            victims AS (
                SELECT id FROM ranked WHERE rn > 1
            ),
            del_translations AS (
                DELETE FROM destination_translations
                WHERE "destinationId" IN (SELECT id FROM victims)
                RETURNING 1
            ),
            del_reviews AS (
                DELETE FROM reviews
                WHERE "destinationId" IN (SELECT id FROM victims)
                RETURNING 1
            ),
            del_favorites AS (
                DELETE FROM favorites
                WHERE "destinationId" IN (SELECT id FROM victims)
                RETURNING 1
            ),
            del_clicks AS (
                DELETE FROM clicks
                WHERE "destinationId" IN (SELECT id FROM victims)
                RETURNING 1
            ),
            del_tours AS (
                DELETE FROM tour_destinations
                WHERE "destinationId" IN (SELECT id FROM victims)
                RETURNING 1
            )
            DELETE FROM destinations
            WHERE id IN (SELECT id FROM victims);
            """
        )
        deleted = int(cur.rowcount or 0)
        conn.commit()
        cur.execute("SELECT COUNT(*) FROM destinations;")
        after = int(cur.fetchone()[0])
        print(f"[OK] Mukerrer kayit silindi: {deleted} (once {before}, sonra {after})")
        return deleted
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()


def copy_translations_into_descriptions(conn) -> int:
    """Cevirideki metni panelin okudugu description alanina tasir."""
    cur = conn.cursor()
    try:
        cur.execute(
            """
            UPDATE destinations d
               SET description = t.description,
                   "shortDescription" = LEFT(t.description, 500),
                   "updatedAt" = NOW()
              FROM destination_translations t
             WHERE t."destinationId" = d.id
               AND t."langCode" = 'tr'
               AND t.description IS NOT NULL
               AND btrim(t.description) <> ''
               AND (d.description IS NULL OR btrim(d.description) = '');
            """
        )
        updated = int(cur.rowcount or 0)
        conn.commit()
        print(f"[OK] Aciklama alani dolduruldu: {updated}")
        return updated
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()


def normalize_stored_image_paths(conn) -> int:
    """/images/yol bicimini masaustu uygulamanin actigi images/yol bicimine cevirir."""
    cur = conn.cursor()
    try:
        cur.execute(
            """
            UPDATE destinations
               SET "thumbnailUrl" = NULLIF(regexp_replace("thumbnailUrl", '^/+', ''), ''),
                   images = ARRAY(
                       SELECT regexp_replace(img, '^/+', '')
                         FROM unnest(COALESCE(images, ARRAY[]::text[])) AS img
                        WHERE img IS NOT NULL AND btrim(img) <> ''
                   )
             WHERE COALESCE("thumbnailUrl", '') LIKE '/%'
                OR EXISTS (
                    SELECT 1 FROM unnest(COALESCE(images, ARRAY[]::text[])) AS img
                     WHERE img LIKE '/%'
                );
            """
        )
        updated = int(cur.rowcount or 0)
        cur.execute(
            """
            UPDATE destinations
               SET "thumbnailUrl" = images[1]
             WHERE ("thumbnailUrl" IS NULL OR btrim("thumbnailUrl") = '')
               AND cardinality(images) >= 1;
            """
        )
        conn.commit()
        print(f"[OK] Resim yolu duzeltildi: {updated}")
        return updated
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()


def find_image_file(dest_folder: Path) -> Optional[str]:
    for ext in ['.jpg', '.jpeg', '.png', '.jfif']:
        images = list(dest_folder.glob(f'*{ext}'))
        if images: return images[0].name
    return None


def normalize_name(name: str) -> str:
    name = re.sub(r'\s*\([^)]+\)\s*$', '', name).strip()
    return re.sub(r'\s+', ' ', name).strip()


def copy_images_to_images_folder(source_root: Path, dest_images: Path):
    dest_images.mkdir(parents=True, exist_ok=True)
    allowed = {il["ad"].lower() for il in ILLER_DATA}
    copied = 0
    for province_dir in source_root.iterdir():
        if not province_dir.is_dir() or province_dir.name.lower() not in allowed:
            continue
        for ext in ['.jpg', '.jpeg', '.png', '.jfif']:
            for img_file in province_dir.rglob(f'*{ext}'):
                if img_file.name.startswith('~$'): continue
                dest_file = dest_images / img_file.name
                if not dest_file.exists():
                    try: shutil.copy2(img_file, dest_file); copied += 1
                    except Exception as e: print(f"  [UYARI] Gorsel kopyalama hatasi: {e}")
    print(f"\n[OK] {copied} gorsel images/ klasorune kopyalandi.")


def get_db_connection():
    return pgdb.connect(
        user=DB_CONFIG["user"], password=DB_CONFIG["password"] or None,
        host=DB_CONFIG["host"], port=int(DB_CONFIG["port"]),
        database=DB_CONFIG["database"], timeout=10,
        application_name="TurizmVeriImportNew",
    )


def resolve_odev_root() -> Path:
    """ödev kökü: prisma -> database -> packages -> turizm-platformu -> ödev."""
    env_root = os.environ.get("ODEV_ROOT")
    if env_root:
        return Path(env_root).resolve()
    return Path(__file__).resolve().parents[4]


def main():
    print("=" * 60 + "\nTURIZM DESTINASYON VERI ICE AKTARMA (YENI SEMA)\n" + "=" * 60)
    source_root = resolve_odev_root()
    images_dest = source_root / "destinasyon" / "images"
    try:
        conn = get_db_connection()
        print("[OK] Veritabanina baglanildi.")
    except Exception as e:
        print(f"[HATA] Veritabani baglanti hatasi: {e}"); return 1

    try:
        setup_provinces_and_categories(conn)
        dedupe_destinations(conn)
        copy_translations_into_descriptions(conn)
        normalize_stored_image_paths(conn)
        copy_images_to_images_folder(source_root, images_dest)

        total_count = 0
        for il in ILLER_DATA:
            il_folder = next(
                (d for d in source_root.iterdir() if d.is_dir() and d.name.lower() == il["ad"].lower()),
                None,
            )
            if il_folder and il_folder.name != 'destinasyon':
                province_id = get_province_id(conn, il["ad"])
                total_count += process_province_new(conn, il_folder, il["ad"], province_id, images_dest)
            else:
                print(f"\n[UYARI] {il['ad']} klasoru bulunamadi.")

        print(f"\n{'='*60}\nTOPLAM: {total_count} destinasyon ice aktarildi.\n{'='*60}")
        cur = conn.cursor()
        try:
            cur.execute('SELECT p.name, COUNT(d.id) FROM provinces p LEFT JOIN destinations d ON d."provinceId" = p.id GROUP BY p.name ORDER BY p.name;')
            for row in cur.fetchall(): print(f"  {row[0]}: {row[1]} destinasyon")
            cur.execute('SELECT c.name, COUNT(d.id) FROM destination_types c LEFT JOIN destinations d ON d."typeId" = c.id GROUP BY c.name ORDER BY COUNT(d.id) DESC;')
            print("\nTur dagilimi:")
            for row in cur.fetchall(): print(f"  {row[0]}: {row[1]}")
            cur.execute(
                """
                SELECT COUNT(*) FROM (
                    SELECT 1 FROM destinations
                    GROUP BY "provinceId", lower(btrim(name)), latitude, longitude
                    HAVING COUNT(*) > 1
                ) s;
                """
            )
            print(f"\nKalan mukerrer grup: {cur.fetchone()[0]}")
            cur.execute(
                "SELECT COUNT(*) FROM destinations WHERE description IS NULL OR btrim(description) = '';"
            )
            print(f"Aciklamasi bos: {cur.fetchone()[0]}")
            cur.execute(
                """
                SELECT COUNT(*) FROM destinations
                 WHERE "thumbnailUrl" IS NULL OR btrim("thumbnailUrl") = '';
                """
            )
            print(f"Resimsiz: {cur.fetchone()[0]}")
            cur.execute(
                """
                SELECT COUNT(*) FROM destinations
                 WHERE COALESCE("thumbnailUrl", '') LIKE '/%';
                """
            )
            print(f"Eski resim yolu: {cur.fetchone()[0]}")
        finally:
            cur.close()
    except Exception as exc:
        print(f"[HATA] Ice aktarma durdu: {exc}")
        return 1
    finally:
        conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())